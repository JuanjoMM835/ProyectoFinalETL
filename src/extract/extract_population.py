"""Leer el Excel municipal del DANE sin modificar la fuente Bronze."""

import logging
from collections import Counter
from pathlib import Path
from typing import Any, Mapping

import pandas as pd
from openpyxl import load_workbook


LOGGER = logging.getLogger(__name__)


def extraer_poblacion(
    configuracion_fuente: Mapping[str, Any], raiz_proyecto: Path
) -> pd.DataFrame:
    """Devolver la hoja completa, junto con su procedencia en ``attrs``.

    La extraccion conserva las notas y filas vacias leidas. El transformador
    las separara de los registros municipales con una regla explicita.
    """
    if configuracion_fuente.get("file_type") != "xlsx":
        raise ValueError("La fuente de poblacion debe declarar file_type: xlsx.")
    ruta = Path(configuracion_fuente["path"])
    if not ruta.is_absolute():
        ruta = Path(raiz_proyecto) / ruta
    ruta = ruta.resolve()
    if not ruta.is_file():
        raise FileNotFoundError("No existe el Excel de poblacion: {}".format(ruta))

    # Copiar evita modificar el bloque de configuracion que recibimos.
    opciones = dict(configuracion_fuente["read_options"])
    hoja = opciones.get("sheet_name")
    encabezado = opciones.get("header")
    if not isinstance(hoja, str):
        raise ValueError("Poblacion requiere el nombre de una sola hoja Excel.")
    if type(encabezado) is not int or encabezado < 0:
        raise ValueError("El encabezado Excel debe ser un indice entero no negativo.")
    if opciones.get("engine") != "openpyxl" or opciones.get("dtype") != "string":
        raise ValueError("Poblacion requiere engine: openpyxl y dtype: string.")
    if opciones.get("keep_default_na") is not False:
        raise ValueError("Declare los vacios explicitamente con keep_default_na: false.")
    # Estas opciones cambiarian las filas fisicas o eliminarian campos de origen.
    if any(nombre in opciones for nombre in ("skiprows", "skipfooter", "nrows", "usecols", "index_col")):
        raise ValueError("La extraccion debe leer la hoja completa despues del encabezado.")

    # Revisar el encabezado real antes de pandas impide que un renombrado
    # automatico oculte columnas duplicadas. read_only nunca escribe el libro.
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        if hoja not in libro.sheetnames:
            raise ValueError("No existe la hoja de poblacion: {}".format(hoja))
        filas_iniciales = list(libro[hoja].iter_rows(
            min_row=1, max_row=encabezado + 1, values_only=True
        ))
    finally:
        libro.close()
    nombres = list(filas_iniciales[-1])
    repetidos = [nombre for nombre, cantidad in Counter(nombres).items()
                 if nombre is not None and cantidad > 1]
    if repetidos:
        raise ValueError("El Excel repite encabezados: {}".format(repetidos))
    faltantes = sorted(set(configuracion_fuente["required_columns"]) - set(nombres))
    if faltantes:
        raise ValueError("Faltan columnas de poblacion obligatorias: {}".format(faltantes))

    LOGGER.info("Leyendo poblacion: %s, hoja %s", ruta.name, hoja)
    datos = pd.read_excel(ruta, **opciones)
    if datos.empty:
        raise ValueError("La hoja de poblacion no contiene registros.")
    # pandas usa indice 0 para header; las filas visibles de Excel empiezan en 1.
    # Conservamos los titulos previos y la ubicacion del encabezado como evidencia.
    datos.attrs["origen_excel"] = {
        "archivo": str(ruta), "hoja": hoja, "header": encabezado,
        "fila_encabezado_excel": encabezado + 1,
        "preambulo": [
            {"fila_excel": numero, "valores": [None if valor is None else str(valor)
                                                for valor in fila]}
            for numero, fila in enumerate(filas_iniciales[:-1], 1)
            if any(valor is not None for valor in fila)
        ],
    }
    LOGGER.info("Extraidas %s filas y %s columnas de poblacion", len(datos), len(datos.columns))
    return datos
